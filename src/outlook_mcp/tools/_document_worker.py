"""Bounded PDF/XLSX extraction, run as a separate process without server credentials.

No OCR or formula evaluation: absent text and formula caches are reported explicitly.
A clean-environment subprocess for crash and resource containment, not a sandbox.
No server package is imported when this file runs as a script. The parent enforces
RSS and wall limits on every platform; OS resource limits are best-effort extras.
"""

from __future__ import annotations

import json
import sys
from contextlib import ExitStack
from io import BytesIO
from zipfile import ZipFile

MAX_CHARS = 1_000_000
MAX_CELLS = 100_000


class RefusalError(Exception):
    """Only controlled, document-independent refusal text may cross the process boundary."""


def resource_limits() -> None:
    try:
        import resource
    except ImportError:
        return
    limits = [(resource.RLIMIT_CPU, 20)]
    if sys.platform != "darwin":
        limits.append((resource.RLIMIT_AS, 256 * 2**20))
    for kind, value in limits:
        try:
            resource.setrlimit(kind, (value, value))
        except (ValueError, OSError):
            pass  # The parent's RSS watchdog and wall timeout remain mandatory.


def append(parts: list[str], text: str, length: int) -> int:
    length += len(text) + bool(parts)
    if length > MAX_CHARS:
        raise RefusalError(
            "The document exceeds one million extracted characters. "
            "Download it for local extraction."
        )
    parts.append(text)
    return length


def pdf(data: bytes) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise RefusalError("This PDF needs a password. Download and unlock it before reading.")
    if len(reader.pages) > 200:
        raise RefusalError("This PDF exceeds 200 pages. Download it for local extraction.")
    parts, empty, length = [], [], 0
    for number, page in enumerate(reader.pages, 1):
        text = page.extract_text(extraction_mode="layout", layout_mode_strip_rotated=False) or ""
        length = append(parts, f"[Page {number}]\n{text}", length)
        if not text.strip():
            empty.append(number)
    if len(empty) == len(reader.pages):
        raise RefusalError(
            "This PDF has no extractable text. It may be scanned: download it and use OCR."
        )
    warnings = [f"Pages without extractable text (blank or need OCR): {empty}"] if empty else []
    return {"text": "\n".join(parts), "pages": len(reader.pages), "warnings": warnings}


def xlsx(data: bytes) -> dict:
    from openpyxl import load_workbook

    with ZipFile(BytesIO(data)) as archive:
        if (
            len(archive.infolist()) > 2000
            or sum(i.file_size for i in archive.infolist()) > 50 * 2**20
        ):
            raise RefusalError(
                "The spreadsheet expands beyond the reading limit. "
                "Download it for local extraction."
            )
    with ExitStack() as stack:
        formulas = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        stack.callback(formulas.close)
        values = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
        stack.callback(values.close)
        parts, sheets, count, length = [], [], 0, 0
        for sheet in formulas:
            sheet.reset_dimensions()  # a stale declared range must not hide populated cells
            cached = values[sheet.title]
            cached.reset_dimensions()
            sheets.append({"name": sheet.title, "state": sheet.sheet_state})
            length = append(
                parts, f"[Sheet {json.dumps(sheet.title)}; {sheet.sheet_state}]", length
            )
            for row, saved in zip(sheet.iter_rows(), cached.iter_rows(), strict=True):
                count += len(row)
                if count > MAX_CELLS:
                    raise RefusalError(
                        "The spreadsheet exceeds 100,000 visited cells. "
                        "Download it for local extraction."
                    )
                for cell, value in zip(row, saved, strict=True):
                    if cell.value is None:
                        continue
                    rendered = json.dumps(cell.value, ensure_ascii=False, default=str)
                    if cell.data_type == "f":
                        rendered += " | saved_result=" + json.dumps(
                            value.value, ensure_ascii=False, default=str
                        )
                    length = append(
                        parts,
                        f"{cell.coordinate}: {rendered} | "
                        f"number_format={json.dumps(cell.number_format)}",
                        length,
                    )
        return {
            "text": "\n".join(parts),
            "sheets": sheets,
            "warnings": [
                "Formulas are not recalculated. saved_result is Excel's stored value; "
                "null means no saved result. "
                "Charts, images and formatting beyond number formats are not extracted."
            ],
        }


def main() -> None:
    resource_limits()
    try:
        # Import checks are separate from parser execution: a parser's own
        # ImportError or UnicodeError is a parse failure, not an install hint.
        if sys.argv[1] == "pdf":
            import pypdf  # noqa: F401
        elif sys.argv[1] == "xlsx":
            import openpyxl  # noqa: F401
        else:
            raise RefusalError("Unsupported document type.")
    except ImportError:
        print(
            json.dumps(
                {
                    "error": "PDF/XLSX reading needs the optional documents extra. "
                    "Ask the operator to install outlook-graph-mcp[documents] "
                    "and restart the server."
                }
            )
        )
        return
    try:
        data = sys.stdin.buffer.read(10 * 2**20 + 1)
        if len(data) > 10 * 2**20:
            raise RefusalError("The attachment exceeds 10 MiB.")
        result = {"pdf": pdf, "xlsx": xlsx}[sys.argv[1]](data)
    except MemoryError:
        result = {"error": "The document parser exceeded its memory limit."}
    except RefusalError as exc:
        result = {"error": str(exc)}
    except Exception as exc:
        result = {
            "error": f"Could not extract this document ({type(exc).__name__}). "
            "Download the original."
        }
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
